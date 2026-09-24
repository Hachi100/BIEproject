"""PTAB synthétiques reproduisant des mises en page rencontrées dans les ministères.

Ils servent de jeu d'essai en attendant les PTAB réels : chaque canevas réel reçu doit
donner lieu à un test (anonymisé si nécessaire) qui fige son interprétation.
"""

from __future__ import annotations

import io
from datetime import date

import openpyxl


def _save(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def ptab_ministere_classique() -> bytes:
    """Titres, en-tête sur deux lignes (chronogramme T1..T4), programme et action fusionnés
    verticalement, une tâche par ligne, sous-total et total général, montants en FCFA."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PTA 2026"
    ws["A1"] = "MINISTÈRE DE LA SANTÉ"
    ws["A2"] = "PLAN DE TRAVAIL ANNUEL BUDGÉTISÉ — GESTION 2026"
    ws.append([])
    ws.append(
        [
            "Programme",
            "Action",
            "Activités",
            "Tâches",
            "Structure responsable",
            "Résultats attendus",
            "Chronogramme",
            None,
            None,
            None,
            "Coût total (FCFA)",
            "Source de financement",
            "Observations",
        ]
    )
    ws.append([None, None, None, None, None, None, "T1", "T2", "T3", "T4", None, None, None])
    ws.merge_cells("G4:J4")
    for col in "ABCDEFKLM":
        ws.merge_cells(f"{col}4:{col}5")
    rows = [
        [
            "Programme 045 : Pilotage et soutien aux services du MS",
            "Action 1 : Gestion des ressources humaines",
            "Organiser la revue à mi-parcours du PTA",
            "Préparer les termes de référence",
            "DPAF",
            "Le taux d'exécution à mi-parcours est connu",
            None,
            "X",
            None,
            None,
            150000,
            "BN",
            None,
        ],
        [
            None,
            None,
            None,
            "Tenir l'atelier de revue (60 participants × 3 jours)",
            "SPSE",
            None,
            None,
            None,
            "X",
            None,
            4500000,
            "BN",
            None,
        ],
        [None, None, None, "Diffuser le rapport de revue", "SPSE", None, None, None, "X", None, 350000, "BN", None],
        [
            None,
            None,
            "Former les gestionnaires de crédits au SIGFP",
            "Recruter un consultant",
            "DPAF",
            "80 gestionnaires formés",
            "X",
            None,
            None,
            None,
            "#REF!",
            "FINEX",
            "Coût à confirmer",
        ],
        [
            None,
            None,
            None,
            "Organiser deux sessions de formation",
            "DPAF",
            None,
            None,
            "X",
            None,
            None,
            "2 000 000",
            "Banque mondiale",
            None,
        ],
        ["Sous-total Action 1", None, None, None, None, None, None, None, None, None, 7000000, None, None],
        [None, None, None, None, None, None, None, None, None, None, None, None, None],
        [None, None, None, None, None, None, None, None, None, None, 999, None, None],
        ["TOTAL GÉNÉRAL", None, None, None, None, None, None, None, None, None, 7000000, None, None],
    ]
    for row in rows:
        ws.append(row)
    ws.merge_cells("A6:A10")
    ws.merge_cells("B6:B10")
    ws.merge_cells("C6:C8")
    ws.merge_cells("F6:F8")
    ws.merge_cells("C9:C10")
    return _save(wb)


def ptab_codes_hierarchiques(montant_programme: int = 10250) -> bytes:
    """Une colonne de codes hiérarchiques, un libellé unique, montants en milliers de FCFA
    ventilés BN / FINEX, dates de début et de fin."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PTAB"
    ws["A1"] = "PTAB 2026 — Ministère de l'Économie et des Finances"
    ws["A2"] = "(montants en milliers de FCFA)"
    ws.append(
        [
            "N°",
            "Libellé",
            "Responsable",
            "Date début",
            "Date fin",
            "Quantité",
            "Prix unitaire",
            "Montant",
            "BN",
            "FINEX",
        ]
    )
    rows = [
        ["026", "Pilotage et soutien aux services du MEF", None, None, None, None, None, montant_programme, None, None],
        ["026.1", "Coordination administrative", None, None, None, None, None, None, None, None],
        [
            "026.1.01",
            "Organiser les conférences budgétaires",
            "DGB",
            date(2026, 7, 1),
            date(2026, 8, 31),
            None,
            None,
            1250,
            1250,
            None,
        ],
        ["026.1.01.1", "Préparer les dossiers", "DPSE", date(2026, 7, 1), date(2026, 7, 15), 10, 25, None, 250, None],
        ["026.1.01.2", "Tenir les séances", "DPSE", date(2026, 8, 1), date(2026, 8, 31), None, None, 1000, 1000, None],
        ["026.1.02", "Moderniser l'archivage", "DSI", date(2026, 3, 1), date(2026, 2, 1), None, None, 9000, 3000, 6000],
        ["026.1.02.1", "Acquérir des serveurs", "DSI", None, None, 2, 4500, 9000, 3000, 6000],
    ]
    for row in rows:
        ws.append(row)
    return _save(wb)


def ptab_csv_francais() -> bytes:
    """Export CSV séparé par des points-virgules, nombres et dates au format français."""
    lines = [
        "Code activité;Activité;Tâche;Responsable;Période;Coût (FCFA);Financement;Mode d'exécution",
        "A1;Campagne de vaccination;Acheter les vaccins;DNSP;T1 - T2;12 500 000;Budget national;Marché",
        ";;Mobiliser les relais communautaires;ZS;T2;1.500.000;Budget national;Régie",
        "A2;Supervision des formations sanitaires;;DDS;Mars à juin;3 250 000,50;UNICEF;Régie",
        ";;;;;850 000;;",
        "TOTAL;;;;;17 600 000,50;;",
    ]
    return "\n".join(lines).encode("cp1252")
