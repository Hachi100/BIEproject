"""Lecture brute des classeurs PTAB (xlsx, xls, csv) en grilles de valeurs.

Les cellules fusionnées sont conservées sous forme de plages : le pipeline décide, selon le
type de colonne, s'il faut propager la valeur (libellés d'un programme fusionnés sur ses
activités) ou non (un montant fusionné ne doit pas être compté plusieurs fois).
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import xlrd


@dataclass
class SheetData:
    name: str
    rows: list[list[object]]
    merged: list[tuple[int, int, int, int]] = field(
        default_factory=list
    )  # (ligne1, col1, ligne2, col2), 0-based inclus
    hidden_rows: set[int] = field(default_factory=set)

    def __post_init__(self) -> None:
        width = max((len(r) for r in self.rows), default=0)
        self.rows = [list(r) + [None] * (width - len(r)) for r in self.rows]
        self._origin: dict[tuple[int, int], tuple[int, int]] = {}
        for r1, c1, r2, c2 in self.merged:
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    if (r, c) != (r1, c1):
                        self._origin[(r, c)] = (r1, c1)

    @property
    def width(self) -> int:
        return len(self.rows[0]) if self.rows else 0

    def value(self, row: int, column: int, fill_merged: bool) -> object:
        """Valeur d'une cellule ; pour une cellule fusionnée non-origine, la valeur d'origine si demandé."""
        origin = self._origin.get((row, column))
        if origin is not None:
            return self.rows[origin[0]][origin[1]] if fill_merged else None
        if row >= len(self.rows) or column >= len(self.rows[row]):
            return None
        return self.rows[row][column]

    def is_merge_continuation(self, row: int, column: int) -> bool:
        return (row, column) in self._origin

    def merge_origin_row(self, row: int, column: int) -> int:
        return self._origin.get((row, column), (row, column))[0]

    def row_values(self, row: int, fill_merged: bool = True) -> list[object]:
        return [self.value(row, c, fill_merged) for c in range(self.width)]


def read_workbook(content: bytes, filename: str) -> list[SheetData]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(content)
    if suffix == ".xls":
        return _read_xls(content)
    if suffix in {".csv", ".txt"}:
        return [_read_csv(content, Path(filename).stem)]
    raise ValueError(f"Format non pris en charge : {suffix or 'inconnu'} (attendu : .xlsx, .xls ou .csv)")


def _read_xlsx(content: bytes) -> list[SheetData]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            continue
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        merged = [(m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1) for m in sheet.merged_cells.ranges]
        hidden = {index - 1 for index, dim in sheet.row_dimensions.items() if dim.hidden}
        sheets.append(SheetData(sheet.title, rows, merged, hidden))
    return sheets


def _read_xls(content: bytes) -> list[SheetData]:
    workbook = xlrd.open_workbook(file_contents=content, formatting_info=True)
    sheets = []
    for sheet in workbook.sheets():
        if sheet.visibility != 0:
            continue
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode))
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(xlrd.error_text_from_code.get(cell.value, "#ERR"))
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        merged = [(r1, c1, r2 - 1, c2 - 1) for r1, r2, c1, c2 in sheet.merged_cells]
        hidden = {r for r in range(sheet.nrows) if sheet.rowinfo_map.get(r) and sheet.rowinfo_map[r].hidden}
        sheets.append(SheetData(sheet.name, rows, merged, hidden))
    return sheets


def _read_csv(content: bytes, name: str) -> SheetData:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows = [
        [cell if cell != "" else None for cell in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)
    ]
    return SheetData(name, rows)
