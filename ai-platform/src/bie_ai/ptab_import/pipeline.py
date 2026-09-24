"""Import intelligent d'un PTAB (Partie 7 §24 « Excel Intelligence ») :

    téléversement → feuilles → en-têtes → correspondance des colonnes → structure
    → valeurs → contrôles → rapprochements → rapport

Principe « zéro perte silencieuse » (Partie 10 §42) : chaque ligne du fichier reçoit un sort
explicite (importée, ignorée avec motif, ou mise en quarantaine avec motif). Aucune règle ne
s'appuie sur un modèle de langage : l'import est déterministe et reproductible. L'IA pourra
proposer des correspondances pour les colonnes non reconnues, sous validation humaine.
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from decimal import Decimal

from openpyxl.utils import get_column_letter

from .columns import AMOUNT_FIELDS, ColumnMatch, combine_header_rows, map_columns, unit_multiplier
from .models import (
    ColumnMapping,
    Discrepancy,
    Funding,
    ImportedActivity,
    ImportedTask,
    ImportReport,
    Level,
    Period,
    RowOutcome,
    RowReport,
    SheetCandidate,
    Summary,
)
from .reader import SheetData, read_workbook
from .values import (
    clean_text,
    is_blank,
    is_schedule_mark,
    month_of,
    normalize,
    parse_amount,
    parse_date,
)

MAX_HEADER_SCAN = 40
LABEL_FIELDS = {"activity_label", "task_label", "generic_label", "program_label", "action_label"}
RECOMMENDED_FIELDS = ["responsible_unit", "total_amount", "period_or_schedule", "funding_source_or_amounts"]
FUNDING_AMOUNT_FIELDS = {
    "amount_national_budget": "national_budget",
    "amount_external": "external",
    "amount_grants": "grants",
    "amount_loans": "loans",
    "amount_counterpart": "counterpart",
}
TOTAL_PREFIXES = ("total", "sous total", "s total", "montant total", "cout total", "grand total", "totaux")
# Un préfixe n'est retenu que s'il est suivi d'un numéro : « Action de sensibilisation » est une activité.
LEVEL_PREFIXES = [
    (Level.PROGRAM, re.compile(r"^(programme|prog)( budgetaire)? (n )?\d")),
    (Level.ACTION, re.compile(r"^action (n )?\d")),
    (Level.ACTIVITY, re.compile(r"^(activite|act) (n )?\d")),
    (Level.TASK, re.compile(r"^(tache|sous activite) (n )?\d")),
]
BANNER_PREFIXES = [(Level.PROGRAM, re.compile(r"^programme\b")), (Level.ACTION, re.compile(r"^action\b"))]
QUARTERLY_AMOUNT_MINIMUM = Decimal(100)  # en dessous : simple coche de chronogramme


class PtabImportError(ValueError):
    """Le fichier ne peut pas être interprété comme un PTAB (motif explicite pour l'utilisateur)."""


@dataclass
class HeaderDetection:
    rows: list[int]
    headers: list[object]
    matches: list[ColumnMatch]
    unmapped: list[int]

    @property
    def score(self) -> int:
        fields = {m.field for m in self.matches}
        has_label = bool(fields & LABEL_FIELDS)
        return len(fields) + (3 if has_label else 0) if has_label else 0

    @property
    def last_row(self) -> int:
        return self.rows[-1]


@dataclass
class _Section:
    code: str | None = None
    label: str | None = None

    def key(self) -> str | None:
        return normalize(self.code) if self.code else (normalize(self.label) if self.label else None)


@dataclass
class _DeclaredTotal:
    row: int
    label: str
    amount: Decimal
    start: int  # index de la première activité couverte
    end: int  # index (exclu) de la dernière activité couverte


@dataclass
class _SectionMark:
    level: Level
    row: int
    label: str
    start: int
    declared: Decimal | None = None


@dataclass
class _State:
    program: _Section = field(default_factory=_Section)
    action: _Section = field(default_factory=_Section)
    activity: ImportedActivity | None = None
    activity_key: str | None = None
    activity_rows: list[int] = field(default_factory=list)


# ---------------------------------------------------------------------------------------
# Entrée principale
# ---------------------------------------------------------------------------------------


def import_ptab(
    content: bytes, filename: str, sheet_name: str | None = None, fiscal_year: int | None = None
) -> ImportReport:
    try:
        sheets = read_workbook(content, filename)
    except ValueError:
        raise
    except Exception as error:  # fichier corrompu, protégé par mot de passe…
        raise PtabImportError(f"Fichier illisible : {error}") from error
    if not sheets:
        raise PtabImportError("Le classeur ne contient aucune feuille visible.")

    detections = {s.name: detect_header(s) for s in sheets}
    candidates = [
        SheetCandidate(
            name=s.name,
            header_row=(detections[s.name].rows[0] + 1) if detections[s.name] else None,
            mapped_columns=len(detections[s.name].matches) if detections[s.name] else 0,
        )
        for s in sheets
    ]
    if sheet_name:
        sheet = next((s for s in sheets if s.name == sheet_name), None)
        if sheet is None:
            raise PtabImportError(
                f"Feuille « {sheet_name} » introuvable. Feuilles : {', '.join(s.name for s in sheets)}."
            )
    else:
        sheet = max(sheets, key=lambda s: detections[s.name].score if detections[s.name] else -1)
    detection = detections[sheet.name]
    if detection is None or detection.score == 0:
        raise PtabImportError(
            f"Aucune ligne d'en-tête reconnue dans la feuille « {sheet.name} » : "
            "les colonnes de libellé (activités, tâches, libellé/désignation) n'ont pas été trouvées."
        )
    return _Importer(sheet, detection, filename, candidates, fiscal_year).run()


def detect_header(sheet: SheetData) -> HeaderDetection | None:
    best: HeaderDetection | None = None
    for r in range(min(MAX_HEADER_SCAN, len(sheet.rows))):
        row = sheet.row_values(r, fill_merged=True)
        if sum(not is_blank(v) for v in row) < 2:
            continue
        matches, unmapped = map_columns(row)
        candidate = HeaderDetection([r], row, matches, unmapped)
        if r + 1 < len(sheet.rows):
            below = sheet.row_values(r + 1, fill_merged=True)
            numeric = sum(isinstance(v, (int, float)) and not isinstance(v, bool) for v in below)
            if numeric <= 1 and sum(not is_blank(v) for v in below) >= 2:
                combined = combine_header_rows(row, below)
                matches2, unmapped2 = map_columns(combined)
                two_rows = HeaderDetection([r, r + 1], combined, matches2, unmapped2)
                if two_rows.score > candidate.score:
                    candidate = two_rows
        if candidate.score >= 5 and (best is None or candidate.score > best.score):
            best = candidate
    return best


# ---------------------------------------------------------------------------------------
# Import d'une feuille
# ---------------------------------------------------------------------------------------


class _Importer:
    def __init__(
        self,
        sheet: SheetData,
        detection: HeaderDetection,
        filename: str,
        candidates: list[SheetCandidate],
        fiscal_year: int | None,
    ):
        self.sheet = sheet
        self.detection = detection
        self.filename = filename
        self.candidates = candidates
        self.title_rows = list(range(detection.rows[0]))
        self.default_multiplier = self._detect_default_multiplier()
        self.matches = self._apply_default_multiplier(detection.matches)
        self.fields: dict[str, ColumnMatch] = {
            m.field: m for m in self.matches if not m.field.startswith(("quarter_", "month_"))
        }
        self._resolve_label_roles()
        self.quarter_columns = {int(m.field.split("_")[1]): m for m in self.matches if m.field.startswith("quarter_")}
        self.month_columns = {int(m.field.split("_")[1]): m for m in self.matches if m.field.startswith("month_")}
        self.fiscal_year = fiscal_year or self._detect_fiscal_year()
        self.structure_mode = "EXPLICIT" if {"activity_label", "task_label"} & self.fields.keys() else "CODED"
        self.rows: list[RowReport] = []
        self.activities: list[ImportedActivity] = []
        self.discrepancies: list[Discrepancy] = []
        self.warnings: list[str] = []
        self.coded_levels: dict[int, Level] = {}
        self.sections: list[_SectionMark] = []
        self.header_values = [normalize(v) if not is_blank(v) else "" for v in detection.headers]

    # -- préparation ------------------------------------------------------------------

    def _detect_default_multiplier(self) -> Decimal:
        texts = [v for r in self.title_rows for v in self.sheet.row_values(r) if not is_blank(v)]
        texts += [v for v in self.detection.headers if not is_blank(v)]
        for text in texts:
            multiplier = unit_multiplier(text)
            words = normalize(text).split()
            if multiplier and ("en" in words or "unite" in words or "(" in str(text)):
                return multiplier
        return Decimal(1)

    def _apply_default_multiplier(self, matches: list[ColumnMatch]) -> list[ColumnMatch]:
        if self.default_multiplier == 1:
            return matches
        adjusted = []
        for m in matches:
            if m.field in AMOUNT_FIELDS and m.multiplier == 1 and unit_multiplier(m.header) is None:
                m = ColumnMatch(m.column, m.header, m.field, m.score, m.method, self.default_multiplier)
            adjusted.append(m)
        return adjusted

    def _resolve_label_roles(self) -> None:
        """Un libellé générique joue le rôle du niveau manquant (tâche ou activité)."""
        generic = self.fields.get("generic_label")
        if generic is None:
            return
        if "activity_label" in self.fields and "task_label" not in self.fields:
            self.fields["task_label"] = generic
            del self.fields["generic_label"]
        elif "task_label" in self.fields and "activity_label" not in self.fields:
            self.fields["activity_label"] = generic
            del self.fields["generic_label"]

    def _detect_fiscal_year(self) -> int | None:
        years: Counter[int] = Counter()
        for r in self.title_rows:
            for value in self.sheet.row_values(r):
                for year in re.findall(r"\b(20[2-4]\d)\b", str(value or "")):
                    years[int(year)] += 1
        return years.most_common(1)[0][0] if years else None

    # -- boucle principale ---------------------------------------------------------------

    def run(self) -> ImportReport:
        for r in self.title_rows:
            outcome = RowOutcome.IGNORED_EMPTY if self._blank_row(r) else RowOutcome.IGNORED_TITLE
            self.rows.append(RowReport(row=r + 1, outcome=outcome))
        for r in self.detection.rows:
            self.rows.append(RowReport(row=r + 1, outcome=RowOutcome.IGNORED_HEADER))

        data_rows = list(range(self.detection.last_row + 1, len(self.sheet.rows)))
        if self.structure_mode == "CODED":
            self.coded_levels = self._infer_coded_levels(data_rows)

        state = _State()
        totals: list[_DeclaredTotal] = []
        subtotal_start = 0
        for r in data_rows:
            report = self._process_row(r, state)
            if report.outcome == RowOutcome.IGNORED_TOTAL:
                declared = self._row_total_amount(r)
                if declared is not None:
                    totals.append(
                        _DeclaredTotal(r, self._first_text(r) or "", declared, subtotal_start, len(self.activities))
                    )
                subtotal_start = len(self.activities)
            if r in self.sheet.hidden_rows:
                report.messages.append("Ligne masquée dans le fichier d'origine.")
            self.rows.append(report)

        self._finalize_activities()
        self._reconcile_sections()
        self._reconcile_totals(totals)
        return self._report()

    def _process_row(self, r: int, state: _State) -> RowReport:
        if self._blank_row(r):
            return RowReport(row=r + 1, outcome=RowOutcome.IGNORED_EMPTY)
        if self._repeated_header(r):
            return RowReport(row=r + 1, outcome=RowOutcome.IGNORED_HEADER, messages=["En-tête répété (saut de page)."])
        if self._total_row(r):
            return RowReport(
                row=r + 1,
                outcome=RowOutcome.IGNORED_TOTAL,
                messages=["Ligne de total : utilisée pour le rapprochement."],
            )

        section = self._section_banner(r)
        if section is not None:
            level, text = section
            self._enter_section(state, level, None, text, r)
            self._record_section_total(r)
            return RowReport(row=r + 1, outcome=RowOutcome.IMPORTED, level=level)

        if self.structure_mode == "EXPLICIT":
            return self._process_explicit(r, state)
        return self._process_coded(r, state)

    # -- mode colonnes explicites ------------------------------------------------------------

    def _process_explicit(self, r: int, state: _State) -> RowReport:
        messages: list[str] = []
        program_code, program_label = (
            self._text(r, "program_code", fill=True),
            self._text(r, "program_label", fill=True),
        )
        action_code, action_label = self._text(r, "action_code", fill=True), self._text(r, "action_label", fill=True)
        activity_code, activity_label = (
            self._text(r, "activity_code", fill=True),
            self._text(r, "activity_label", fill=True),
        )
        task_label = self._text(r, "task_label", fill=False)

        if program_code or program_label:
            self._enter_section(state, Level.PROGRAM, program_code, program_label, r, keep_if_same=True)
        if action_code or action_label:
            self._enter_section(state, Level.ACTION, action_code, action_label, r, keep_if_same=True)

        level: Level | None = None
        if activity_label or activity_code:
            key = normalize(activity_code) if activity_code else normalize(activity_label)
            if key != state.activity_key:
                self._start_activity(state, r, activity_code, activity_label or activity_code or "")
                level = Level.ACTIVITY
        if task_label:
            if state.activity is None:
                return RowReport(
                    row=r + 1,
                    outcome=RowOutcome.QUARANTINED,
                    level=Level.TASK,
                    messages=["Tâche sans activité de rattachement : renseigner l'activité dans le fichier."],
                )
            if level == Level.ACTIVITY:
                self._fill_activity_descriptors(state.activity, r)
            self._add_task(state, r, None, task_label, messages)
            return self._row_report(r, Level.TASK, messages)

        if level == Level.ACTIVITY:
            self._fill_activity(state.activity, r, messages, first_row=True)
            return self._row_report(r, Level.ACTIVITY, messages)
        if state.activity is not None and (activity_label or activity_code):
            # ligne supplémentaire de la même activité (ex. une ligne par source de financement)
            self._fill_activity(state.activity, r, messages, first_row=False)
            messages.append(f"Ligne complémentaire de l'activité « {state.activity.label} ».")
            return self._row_report(r, Level.ACTIVITY, messages)
        if program_code or program_label or action_code or action_label:
            return self._section_row_report(r, Level.ACTION if (action_code or action_label) else Level.PROGRAM)
        return self._orphan_values_report(r)

    # -- mode libellé + code hiérarchique -----------------------------------------------------

    def _infer_coded_levels(self, data_rows: list[int]) -> dict[int, Level]:
        depths: dict[int, int] = {}
        levels: dict[int, Level] = {}
        for r in data_rows:
            label = self._text(r, "generic_label", fill=False)
            if not label or self._total_row(r) or self._blank_row(r):
                continue
            prefixed = self._level_from_prefix(label)
            if prefixed:
                levels[r] = prefixed
                continue
            code = self._text(r, "generic_code", fill=False)
            if code:
                depths[r] = len([s for s in re.split(r"[.\-/\s]+", code.strip(" .")) if s])
        distinct = sorted(set(depths.values()))
        order = {1: [Level.ACTIVITY], 2: [Level.ACTIVITY, Level.TASK], 3: [Level.ACTION, Level.ACTIVITY, Level.TASK]}
        ladder = order.get(len(distinct), [Level.PROGRAM, Level.ACTION, Level.ACTIVITY, Level.TASK])
        ladder = [Level.PROGRAM] * (len(distinct) - len(ladder)) + ladder
        by_depth = dict(zip(distinct, ladder, strict=False))
        for r, depth in depths.items():
            levels[r] = by_depth[depth]
        if distinct:
            self.warnings.append(
                "Niveaux déduits de la profondeur des codes : "
                + ", ".join(f"{d} segment(s) → {by_depth[d].value}" for d in distinct)
                + ". À confirmer."
            )
        return levels

    def _process_coded(self, r: int, state: _State) -> RowReport:
        messages: list[str] = []
        label = self._text(r, "generic_label", fill=False)
        code = self._text(r, "generic_code", fill=False)
        if not label:
            return self._orphan_values_report(r)
        level = self.coded_levels.get(r)
        if level is None:
            level = Level.TASK if state.activity is not None else Level.ACTIVITY
            messages.append(f"Ligne sans code : niveau {level.value} déduit de sa position.")
        if level in (Level.PROGRAM, Level.ACTION):
            self._enter_section(state, level, code, label, r)
            return self._section_row_report(r, level, messages)
        if level == Level.ACTIVITY:
            self._start_activity(state, r, code, label)
            self._fill_activity(state.activity, r, messages, first_row=True)
            return self._row_report(r, Level.ACTIVITY, messages)
        if state.activity is None:
            return RowReport(
                row=r + 1,
                outcome=RowOutcome.QUARANTINED,
                level=Level.TASK,
                messages=["Tâche sans activité de rattachement."],
            )
        self._add_task(state, r, code, label, messages)
        return self._row_report(r, Level.TASK, messages)

    # -- construction des objets ---------------------------------------------------------------

    def _enter_section(
        self,
        state: _State,
        level: Level,
        code: str | None,
        label: str | None,
        r: int,
        keep_if_same: bool = False,
    ) -> None:
        section = _Section(code, label)
        current = state.program if level == Level.PROGRAM else state.action
        if keep_if_same and section.key() == current.key():
            return
        self.sections.append(_SectionMark(level, r, " ".join(filter(None, [code, label])), len(self.activities)))
        if level == Level.PROGRAM:
            state.program = section
            state.action = _Section()
        else:
            state.action = section
        state.activity = None
        state.activity_key = None

    def _start_activity(self, state: _State, r: int, code: str | None, label: str) -> None:
        activity = ImportedActivity(
            source_row=r + 1,
            program_code=_program_code(state.program),
            program_label=_strip_prefix(state.program.label),
            action_code=state.action.code,
            action_label=_strip_prefix(state.action.label),
            code=code,
            label=_strip_prefix(label) or label,
        )
        self.activities.append(activity)
        state.activity = activity
        state.activity_key = normalize(code) if code else normalize(label)

    def _add_task(self, state: _State, r: int, code: str | None, label: str, messages: list[str]) -> None:
        activity = state.activity
        assert activity is not None
        task = ImportedTask(source_row=r + 1, code=code, label=_strip_prefix(label) or label)
        task.responsible = self._text(r, "responsible_unit", fill=False)
        task.unit = self._text(r, "unit", fill=False)
        task.observations = self._text(r, "observations", fill=False)
        task.economic_nature = self._text(r, "economic_nature", fill=False)
        task.quantity = self._number(r, "quantity", messages)
        task.unit_price = self._amount(r, "unit_price", messages, allow_merged=False)
        task.period = self._period(r, messages)
        task.funding = self._funding(r, messages, allow_merged=False)
        amount = self._amount(r, "total_amount", messages, allow_merged=False)
        task.amount = self._reconcile_row_amount(r, amount, task.quantity, task.unit_price, task.funding, messages)
        activity.tasks.append(task)
        # attributs d'activité fusionnés verticalement sur les lignes des tâches
        self._merge_activity_level_cells(activity, r, messages)

    def _fill_activity(self, activity: ImportedActivity, r: int, messages: list[str], first_row: bool) -> None:
        for attribute, field_name in (
            ("responsible_unit", "responsible_unit"),
            ("associated_units", "associated_units"),
            ("expected_result", "expected_result"),
            ("indicator", "indicator"),
            ("target", "target"),
            ("location", "location"),
            ("economic_nature", "economic_nature"),
            ("observations", "observations"),
        ):
            value = self._text(r, field_name, fill=True)
            current = getattr(activity, attribute)
            if value and not current:
                setattr(activity, attribute, value)
            elif value and current and normalize(value) != normalize(current) and not first_row:
                messages.append(
                    f"{attribute} différent sur une ligne complémentaire : « {value} » (conservé : « {current} »)."
                )
        mode = self._text(r, "execution_mode", fill=True)
        if mode and not activity.execution_mode:
            activity.execution_mode = _execution_mode(mode) or mode
            if not _execution_mode(mode):
                messages.append(f"Mode d'exécution non reconnu : « {mode} ».")
        period = self._period(r, messages)
        activity.period = _merge_periods(activity.period, period)
        funding = self._funding(r, messages, allow_merged=True)
        activity.funding = _merge_funding(activity.funding, funding)
        amount = self._amount(r, "total_amount", messages, allow_merged=True)
        quantity = self._number(r, "quantity", messages)
        unit_price = self._amount(r, "unit_price", messages, allow_merged=True)
        amount = self._reconcile_row_amount(r, amount, quantity, unit_price, funding, messages)
        if amount is not None:
            activity.declared_amount = (activity.declared_amount or Decimal(0)) + amount

    def _fill_activity_descriptors(self, activity: ImportedActivity, r: int) -> None:
        """Descripteurs d'activité portés par la ligne de sa première tâche (hors montants)."""
        for attribute in ("responsible_unit", "associated_units", "expected_result", "indicator", "target", "location"):
            value = self._text(r, attribute, fill=True)
            if value and not getattr(activity, attribute):
                setattr(activity, attribute, value)
        mode = self._text(r, "execution_mode", fill=True)
        if mode and not activity.execution_mode:
            activity.execution_mode = _execution_mode(mode) or mode

    def _merge_activity_level_cells(self, activity: ImportedActivity, r: int, messages: list[str]) -> None:
        """Montants et libellés d'activité fusionnés sur plusieurs lignes de tâches : comptés une seule fois."""
        for field_name in ("responsible_unit", "expected_result", "indicator", "location", "execution_mode"):
            match = self.fields.get(field_name)
            if match and self._vertical_merge_origin(r, match.column) and not getattr(activity, field_name):
                value = clean_text(self.sheet.value(r, match.column, fill_merged=True))
                setattr(
                    activity, field_name, _execution_mode(value) or value if field_name == "execution_mode" else value
                )
        match = self.fields.get("total_amount")
        if match and self._vertical_merge_origin(r, match.column):
            parsed = parse_amount(self.sheet.value(r, match.column, fill_merged=False))
            if parsed.ok and parsed.value is not None:
                activity.declared_amount = (activity.declared_amount or Decimal(0)) + parsed.value * match.multiplier
                messages.append("Montant fusionné sur plusieurs lignes : rattaché à l'activité, compté une seule fois.")

    # -- lecture des cellules --------------------------------------------------------------------

    def _text(self, r: int, field_name: str, fill: bool) -> str | None:
        match = self.fields.get(field_name)
        if match is None:
            return None
        value = clean_text(self.sheet.value(r, match.column, fill_merged=fill))
        if value and value.upper().startswith("#") and value.upper().endswith(("!", "?", "A")):
            return None  # erreur Excel dans un libellé
        return value

    def _vertical_merge_origin(self, r: int, column: int) -> bool:
        for r1, c1, r2, _c2 in self.sheet.merged:
            if r1 == r and c1 <= column <= _c2 and r2 > r1:
                return True
        return False

    def _amount(self, r: int, field_name: str, messages: list[str], allow_merged: bool) -> Decimal | None:
        match = self.fields.get(field_name)
        if match is None:
            return None
        if self.sheet.is_merge_continuation(r, match.column):
            return None
        if not allow_merged and self._vertical_merge_origin(r, match.column):
            return None  # rattaché à l'activité par _merge_activity_level_cells
        parsed = parse_amount(self.sheet.value(r, match.column, fill_merged=False))
        if not parsed.ok:
            messages.append(f"{match.header} : {parsed.error}")
            return None
        return None if parsed.value is None else parsed.value * match.multiplier

    def _number(self, r: int, field_name: str, messages: list[str]) -> Decimal | None:
        match = self.fields.get(field_name)
        if match is None or self.sheet.is_merge_continuation(r, match.column):
            return None
        parsed = parse_amount(self.sheet.value(r, match.column, fill_merged=False))
        if not parsed.ok:
            messages.append(f"{match.header} : {parsed.error}")
            return None
        return parsed.value

    def _funding(self, r: int, messages: list[str], allow_merged: bool) -> Funding:
        funding = Funding()
        for field_name, attribute in FUNDING_AMOUNT_FIELDS.items():
            setattr(funding, attribute, self._amount(r, field_name, messages, allow_merged))
        text = self._text(r, "funding_source", fill=True)
        funding.source_text = text
        amounts = {a: getattr(funding, a) for a in FUNDING_AMOUNT_FIELDS.values() if getattr(funding, a)}
        if len(amounts) > 1:
            funding.source = "MIXED"
        elif amounts:
            funding.source = {
                "national_budget": "NATIONAL_BUDGET",
                "external": "EXTERNAL",
                "grants": "GRANT",
                "loans": "LOAN",
                "counterpart": "COUNTERPART",
            }[next(iter(amounts))]
        elif text:
            funding.source = _funding_source(text)
        return funding

    def _period(self, r: int, messages: list[str]) -> Period:
        period = Period()
        start = self._date(r, "start_date", messages)
        end = self._date(r, "end_date", messages)
        if start and end and end < start:
            messages.append(f"Date de fin ({end}) antérieure à la date de début ({start}) : dates ignorées.")
            start = end = None
        period.start, period.end = start, end
        for quarter, match in sorted(self.quarter_columns.items()):
            value = self.sheet.value(r, match.column, fill_merged=True)
            if is_schedule_mark(value):
                period.quarters.append(quarter)
                parsed = parse_amount(value)
                if parsed.ok and parsed.value is not None and parsed.value >= QUARTERLY_AMOUNT_MINIMUM:
                    period.quarterly_amounts[quarter] = parsed.value * self.default_multiplier
        for month, match in sorted(self.month_columns.items()):
            if is_schedule_mark(self.sheet.value(r, match.column, fill_merged=True)):
                period.months.append(month)
        text = self._text(r, "period", fill=True)
        if text:
            period.text = text
            quarters, months = _parse_period_text(text)
            period.quarters = sorted(set(period.quarters) | set(quarters))
            period.months = sorted(set(period.months) | set(months))
        if self.fiscal_year and start and start.year != self.fiscal_year:
            messages.append(f"Date de début hors de l'exercice {self.fiscal_year} : {start}.")
        return period

    def _date(self, r: int, field_name: str, messages: list[str]):
        match = self.fields.get(field_name)
        if match is None:
            return None
        parsed = parse_date(self.sheet.value(r, match.column, fill_merged=True), self.fiscal_year)
        if not parsed.ok:
            messages.append(f"{match.header} : {parsed.error}")
            return None
        return parsed.value

    # -- contrôles et rapprochements -------------------------------------------------------------

    def _reconcile_row_amount(
        self,
        r: int,
        amount: Decimal | None,
        quantity: Decimal | None,
        unit_price: Decimal | None,
        funding: Funding,
        messages: list[str],
    ) -> Decimal | None:
        if quantity is not None and unit_price is not None:
            computed = quantity * unit_price
            if amount is None:
                amount = computed
                messages.append("Montant calculé : quantité × prix unitaire.")
            elif abs(computed - amount) > max(Decimal(1), amount * Decimal("0.005")):
                self.discrepancies.append(
                    Discrepancy(
                        row=r + 1,
                        kind="QUANTITE_X_PRIX",
                        declared=amount,
                        computed=computed,
                        message=f"Montant saisi {amount} ≠ quantité × prix unitaire {computed}.",
                    )
                )
        parts = [getattr(funding, a) for a in FUNDING_AMOUNT_FIELDS.values() if getattr(funding, a) is not None]
        if parts:
            funded = sum(parts, Decimal(0))
            if amount is None:
                amount = funded
            elif abs(funded - amount) > Decimal(1):
                self.discrepancies.append(
                    Discrepancy(
                        row=r + 1,
                        kind="VENTILATION_FINANCEMENT",
                        declared=amount,
                        computed=funded,
                        message=f"Montant total {amount} ≠ somme des financements {funded}.",
                    )
                )
        return amount

    def _finalize_activities(self) -> None:
        for activity in self.activities:
            task_amounts = [t.amount for t in activity.tasks if t.amount is not None]
            activity.tasks_amount = sum(task_amounts, Decimal(0)) if task_amounts else None
            if activity.tasks_amount is not None and activity.declared_amount is not None:
                if abs(activity.tasks_amount - activity.declared_amount) > Decimal(1):
                    self.discrepancies.append(
                        Discrepancy(
                            row=activity.source_row,
                            kind="ACTIVITE_VS_TACHES",
                            declared=activity.declared_amount,
                            computed=activity.tasks_amount,
                            message=f"Activité « {activity.label} » : montant déclaré {activity.declared_amount} "
                            f"≠ somme des tâches {activity.tasks_amount}. Le détail des tâches est retenu.",
                        )
                    )
            activity.amount = activity.tasks_amount if activity.tasks_amount is not None else activity.declared_amount
            # financement et calendrier de l'activité : agrégés depuis les tâches s'ils ne sont pas saisis
            # sur la ligne de l'activité (jamais les deux, pour ne pas compter deux fois)
            if activity.tasks and not _has_amounts(activity.funding):
                aggregated = Funding()
                for task in activity.tasks:
                    aggregated = _merge_funding(aggregated, task.funding)
                activity.funding = _merge_funding(activity.funding, aggregated)
            if activity.tasks:
                quarterly = dict(activity.period.quarterly_amounts)
                for task in activity.tasks:
                    activity.period = _merge_periods(
                        activity.period, task.period.model_copy(update={"quarterly_amounts": {}})
                    )
                if not quarterly:
                    for task in activity.tasks:
                        for quarter, amount in task.period.quarterly_amounts.items():
                            quarterly[quarter] = quarterly.get(quarter, Decimal(0)) + amount
                activity.period.quarterly_amounts = quarterly

    def _row_total_amount(self, r: int) -> Decimal | None:
        match = self.fields.get("total_amount")
        if match is None:
            return None
        parsed = parse_amount(self.sheet.value(r, match.column, fill_merged=False))
        return parsed.value * match.multiplier if parsed.ok and parsed.value is not None else None

    def _reconcile_sections(self) -> None:
        """Montant saisi sur une ligne de programme ou d'action ↔ somme de ses activités."""
        for index, section in enumerate(self.sections):
            if section.declared is None:
                continue
            end = len(self.activities)
            for following in self.sections[index + 1 :]:
                if following.level == Level.PROGRAM or following.level == section.level:
                    end = following.start
                    break
            computed = self._sum_amounts(self.activities[section.start : end])
            if computed is not None and abs(computed - section.declared) > Decimal(1):
                self.discrepancies.append(
                    Discrepancy(
                        row=section.row + 1,
                        kind=f"{section.level.value}_VS_ACTIVITES",
                        declared=section.declared,
                        computed=computed,
                        message=f"{section.label} : montant déclaré {section.declared} ≠ somme de ses activités {computed}.",
                    )
                )

    def _reconcile_totals(self, totals: list[_DeclaredTotal]) -> None:
        """Lignes de total : le total général couvre tout, un sous-total couvre les activités depuis le total précédent."""
        general = [t for t in totals if "general" in normalize(t.label)]
        if not general and len(totals) == 1 and totals[0].start == 0 and totals[0].end == len(self.activities):
            general = totals
        for total in totals:
            covered = self.activities if total in general else self.activities[total.start : total.end]
            computed = self._sum_amounts(covered)
            if computed is None or abs(computed - total.amount) <= Decimal(1):
                continue
            if total in general:
                self.discrepancies.append(
                    Discrepancy(
                        row=total.row + 1,
                        kind="TOTAL_GENERAL",
                        declared=total.amount,
                        computed=computed,
                        message=f"Total général déclaré {total.amount} ≠ somme des activités importées {computed}.",
                    )
                )
            else:
                self.discrepancies.append(
                    Discrepancy(
                        row=total.row + 1,
                        kind="SOUS_TOTAL",
                        declared=total.amount,
                        computed=computed,
                        message=f"« {total.label} » : {total.amount} déclaré, {computed} calculé sur les activités qui précèdent. À vérifier.",
                    )
                )

    @staticmethod
    def _sum_amounts(activities: list[ImportedActivity]) -> Decimal | None:
        amounts = [(a.tasks_amount if a.tasks_amount is not None else a.declared_amount) for a in activities]
        amounts = [a for a in amounts if a is not None]
        return sum(amounts, Decimal(0)) if amounts else None

    # -- classification de lignes ----------------------------------------------------------------

    def _blank_row(self, r: int) -> bool:
        return all(is_blank(self.sheet.value(r, c, fill_merged=False)) for c in range(self.sheet.width))

    def _first_text(self, r: int) -> str | None:
        for c in range(self.sheet.width):
            value = self.sheet.value(r, c, fill_merged=False)
            if isinstance(value, str) and value.strip():
                return value.strip()
        return None

    def _total_row(self, r: int) -> bool:
        text = self._first_text(r)
        return bool(text) and normalize(text).startswith(TOTAL_PREFIXES)

    def _repeated_header(self, r: int) -> bool:
        values = [normalize(v) if not is_blank(v) else "" for v in self.sheet.row_values(r)]
        mapped = [m.column for m in self.matches if self.header_values[m.column]]
        if len(mapped) < 3:
            return False
        same = sum(1 for c in mapped if c < len(values) and values[c] == self.header_values[c])
        return same >= 0.6 * len(mapped)

    def _section_banner(self, r: int) -> tuple[Level, str] | None:
        """Ligne-titre d'un programme ou d'une action (souvent fusionnée sur toute la largeur)."""
        texts = [(c, self.sheet.value(r, c, fill_merged=False)) for c in range(self.sheet.width)]
        non_blank = [(c, v) for c, v in texts if not is_blank(v)]
        if len(non_blank) != 1 or not isinstance(non_blank[0][1], str):
            return None
        column, text = non_blank[0]
        spans_row = any(r1 == r and c1 == column and c2 > c1 + 1 for r1, c1, _r2, c2 in self.sheet.merged)
        level = self._level_from_prefix(text)
        if level is None and spans_row:
            level = next((lvl for lvl, pattern in BANNER_PREFIXES if pattern.match(normalize(text))), None)
        if level not in (Level.PROGRAM, Level.ACTION):
            return None
        section_columns = {
            self.fields[f].column
            for f in ("program_code", "program_label", "action_code", "action_label")
            if f in self.fields
        }
        label_columns = {
            self.fields[f].column for f in ("activity_label", "task_label", "generic_label") if f in self.fields
        }
        if spans_row or column in section_columns or column not in label_columns:
            return level, text.strip()
        return None

    @staticmethod
    def _level_from_prefix(text: str) -> Level | None:
        normalized = normalize(text)
        for level, pattern in LEVEL_PREFIXES:
            if pattern.match(normalized):
                return level
        return None

    # -- rapports ------------------------------------------------------------------------------

    def _row_report(self, r: int, level: Level, messages: list[str]) -> RowReport:
        outcome = (
            RowOutcome.IMPORTED_WITH_WARNINGS
            if any(not m.startswith(("Montant calculé", "Montant fusionné", "Ligne complémentaire")) for m in messages)
            else RowOutcome.IMPORTED
        )
        return RowReport(row=r + 1, outcome=outcome, level=level, messages=messages)

    def _section_row_report(self, r: int, level: Level, messages: list[str] | None = None) -> RowReport:
        messages = list(messages or [])
        if self._record_section_total(r) is not None:
            messages.append(f"Montant porté par la ligne {level.value} : rapproché de la somme de ses activités.")
        return RowReport(row=r + 1, outcome=RowOutcome.IMPORTED, level=level, messages=messages)

    def _record_section_total(self, r: int) -> Decimal | None:
        amount = self._row_total_amount(r)
        if amount is not None and self.sections and self.sections[-1].row == r:
            self.sections[-1].declared = amount
        return amount

    def _orphan_values_report(self, r: int) -> RowReport:
        return RowReport(
            row=r + 1,
            outcome=RowOutcome.QUARANTINED,
            messages=[
                "Ligne sans libellé d'activité ni de tâche : impossible de la rattacher. Valeurs : "
                + " | ".join(str(v) for v in self.sheet.row_values(r, fill_merged=False) if not is_blank(v))[:300]
            ],
        )

    def _report(self) -> ImportReport:
        counts = Counter(row.outcome for row in self.rows)
        fields = set(self.fields)
        missing = []
        if "responsible_unit" not in fields:
            missing.append("responsible_unit")
        if not ({"total_amount", "unit_price"} | set(FUNDING_AMOUNT_FIELDS)) & fields:
            missing.append("total_amount")
        if not ({"start_date", "end_date", "period"} & fields) and not self.quarter_columns and not self.month_columns:
            missing.append("period_or_schedule")
        if not ({"funding_source"} | set(FUNDING_AMOUNT_FIELDS)) & fields:
            missing.append("funding_source_or_amounts")
        total = self._sum_amounts(self.activities) or Decimal(0)
        return ImportReport(
            file_name=self.filename,
            sheet=self.sheet.name,
            sheets=self.candidates,
            header_rows=[r + 1 for r in self.detection.rows],
            fiscal_year=self.fiscal_year,
            default_multiplier=self.default_multiplier,
            structure_mode=self.structure_mode,
            columns=[
                ColumnMapping(
                    column=get_column_letter(m.column + 1),
                    header=m.header,
                    field=m.field,
                    score=m.score,
                    method=m.method,
                    multiplier=m.multiplier,
                )
                for m in self.matches
            ],
            unmapped_columns=[
                f"{get_column_letter(c + 1)} ({self.detection.headers[c]})" for c in self.detection.unmapped
            ],
            missing_recommended_fields=missing,
            activities=self.activities,
            rows=sorted(self.rows, key=lambda row: row.row),
            discrepancies=self.discrepancies,
            warnings=self.warnings,
            summary=Summary(
                rows_total=len(self.sheet.rows),
                imported=counts[RowOutcome.IMPORTED],
                imported_with_warnings=counts[RowOutcome.IMPORTED_WITH_WARNINGS],
                ignored={o.value: counts[o] for o in RowOutcome if o.value.startswith("IGNORED") and counts[o]},
                quarantined=counts[RowOutcome.QUARANTINED],
                activities=len(self.activities),
                tasks=sum(len(a.tasks) for a in self.activities),
                total_amount=total,
            ),
        )


# ---------------------------------------------------------------------------------------
# Normalisations métier
# ---------------------------------------------------------------------------------------


def _strip_prefix(label: str | None) -> str | None:
    """« Activité 1.2 : Former les agents » → « Former les agents » (le code reste dans `code`)."""
    if not label:
        return label
    stripped = re.sub(
        r"^\s*(programme|prog|action|activit[eé]|t[aâ]che|sous[- ]activit[eé])\s*(n[°o]?\s*)?[\w.\-]*\s*[:\-–]\s*",
        "",
        label,
        flags=re.IGNORECASE,
    )
    return stripped.strip() or label


def _program_code(section: _Section) -> str | None:
    for candidate in (section.code, section.label):
        if candidate:
            match = re.search(r"\b(\d{2,4})\b", candidate)
            if match:
                digits = match.group(1).lstrip("0") or "0"
                if len(digits) <= 3:
                    return digits.zfill(3)
    return None


def _execution_mode(text: str | None) -> str | None:
    t = normalize(text or "")
    if not t:
        return None
    if "mixte" in t:
        return "MIXED"
    if any(w in t for w in ("regie", "interne", "direct", "faire soi")):
        return "DIRECT"
    if any(
        w in t
        for w in (
            "marche",
            "passation",
            "prestataire",
            "externe",
            "faire faire",
            "sous traitance",
            "indirect",
            "consultant",
        )
    ):
        return "INDIRECT"
    return None


def _funding_source(text: str) -> str | None:
    t = normalize(text)
    sources = set()
    if re.search(r"\b(bn|bne|budget national|budget de l etat|tresor|etat|ressources internes)\b", t):
        sources.add("NATIONAL_BUDGET")
    if re.search(r"\b(pret|prets|emprunt)", t):
        sources.add("LOAN")
    if re.search(r"\b(don|dons|subvention)", t):
        sources.add("GRANT")
    if re.search(r"\bcontrepartie", t):
        sources.add("COUNTERPART")
    if re.search(r"\b(finex|ptf|exterieur|bailleur|banque mondiale|bad|afd|ue|unicef|pnud|oms|fmi|bid|boad)\b", t):
        sources.add("EXTERNAL")
    if len(sources) > 1:
        return "MIXED"
    return next(iter(sources), None)


def _parse_period_text(text: str) -> tuple[list[int], list[int]]:
    t = normalize(text)
    quarters = [int(q) for q in re.findall(r"\bt\s*([1-4])\b", t)]
    quarters += [int(q) for q in re.findall(r"\b([1-4])\s*(?:er|e|eme|ieme)?\s*trim", t)]
    if len(quarters) >= 2 and re.search(r"\b(a|au|-)\b|\s-\s", t):
        quarters = list(range(min(quarters), max(quarters) + 1))
    months = [m for m in (month_of(w) for w in t.split()) if m]
    if len(months) >= 2:
        months = list(range(min(months), max(months) + 1))
    return sorted(set(quarters)), sorted(set(months))


def _merge_periods(current: Period, new: Period) -> Period:
    return Period(
        start=min(filter(None, [current.start, new.start]), default=None),
        end=max(filter(None, [current.end, new.end]), default=None),
        quarters=sorted(set(current.quarters) | set(new.quarters)),
        months=sorted(set(current.months) | set(new.months)),
        quarterly_amounts={
            q: current.quarterly_amounts.get(q, Decimal(0)) + new.quarterly_amounts.get(q, Decimal(0))
            for q in set(current.quarterly_amounts) | set(new.quarterly_amounts)
        },
        text=current.text or new.text,
    )


def _has_amounts(funding: Funding) -> bool:
    return any(getattr(funding, attribute) is not None for attribute in FUNDING_AMOUNT_FIELDS.values())


def _merge_funding(current: Funding, new: Funding) -> Funding:
    merged = Funding()
    for attribute in FUNDING_AMOUNT_FIELDS.values():
        values = [v for v in (getattr(current, attribute), getattr(new, attribute)) if v is not None]
        setattr(merged, attribute, sum(values, Decimal(0)) if values else None)
    texts = [t for t in (current.source_text, new.source_text) if t]
    merged.source_text = " ; ".join(dict.fromkeys(texts)) or None
    sources = {s for s in (current.source, new.source) if s}
    merged.source = "MIXED" if len(sources) > 1 else next(iter(sources), None)
    return merged
